"""
Excel fond faylidan (umumiy_fond.xlsx) kitoblarni import qiladi.

Ustunlar (aynan shunday bo'lishi kerak, 1-qatorda sarlavha sifatida):
Id | BBK | UDK | Kitob mualliflari | Kitob nomi | Turi | Tili |
Nashr qilingan yil | Kitob nashriyoti nomi | Narxi

Ishlatish:
    python manage.py import_fond /yo'l/umumiy_fond.xlsx
    python manage.py import_fond /yo'l/umumiy_fond.xlsx --keep-old   (eskisini o'chirmaydi)

Standart holatda (--keep-old berilmasa) avval quyidagilar butunlay
o'chiriladi: Reservation, Loan, BookCopy, Book, Author, Category.
"""
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.books.models import Book, BookCopy
from apps.catalog.models import Author, Category
from apps.loans.models import Loan, Reservation

REQUIRED_HEADERS = [
    "Id", "BBK", "UDK", "Kitob mualliflari", "Kitob nomi", "Turi", "Tili",
    "Nashr qilingan yil", "Kitob nashriyoti nomi", "Narxi",
]


def clean(value):
    """Har qanday katak qiymatini toza matnga aylantiradi."""
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


class Command(BaseCommand):
    help = "umumiy_fond.xlsx fayldan kitoblarni bazaga import qiladi."

    def add_arguments(self, parser):
        parser.add_argument("excel_path", type=str, help="xlsx fayl manzili")
        parser.add_argument(
            "--keep-old", action="store_true",
            help="Eski ma'lumotlarni o'chirmasdan, ustiga qo'shib boradi",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        path = Path(options["excel_path"])
        if not path.exists():
            raise CommandError(f"Fayl topilmadi: {path}")

        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [clean(h) for h in next(rows)]

        missing = [h for h in REQUIRED_HEADERS if h not in headers]
        if missing:
            raise CommandError(f"Quyidagi ustunlar faylda topilmadi: {missing}")

        col = {h: i for i, h in enumerate(headers)}

        if not options["keep_old"]:
            self.stdout.write("Eski ma'lumotlar o'chirilmoqda...")
            Reservation.objects.all().delete()
            Loan.objects.all().delete()
            BookCopy.objects.all().delete()
            Book.objects.all().delete()
            Author.objects.all().delete()
            Category.objects.all().delete()

        created_books = 0
        reused_books = 0
        created_copies = 0
        skipped = 0
        total = 0

        for row in rows:
            total += 1
            title = clean(row[col["Kitob nomi"]])
            if not title:
                skipped += 1
                continue

            author_name = clean(row[col["Kitob mualliflari"]]) or "Noma'lum muallif"
            category_name = clean(row[col["Turi"]]) or "Boshqa"
            old_id = clean(row[col["Id"]])
            bbk = clean(row[col["BBK"]])
            udk = clean(row[col["UDK"]])
            language = clean(row[col["Tili"]])
            publisher = clean(row[col["Kitob nashriyoti nomi"]])
            price = clean(row[col["Narxi"]])

            year_raw = row[col["Nashr qilingan yil"]]
            try:
                year = int(year_raw) if year_raw not in (None, "") else None
            except (TypeError, ValueError):
                year = None

            author, _ = Author.objects.get_or_create(full_name=author_name[:200])
            category, _ = Category.objects.get_or_create(name=category_name[:150])

            desc_bits = []
            if language:
                desc_bits.append(f"Tili: {language}")
            if bbk:
                desc_bits.append(f"BBK: {bbk}")
            if udk:
                desc_bits.append(f"UDK: {udk}")
            if old_id:
                desc_bits.append(f"Fond ID: {old_id}")
            description = " | ".join(desc_bits)

            book, created = Book.objects.get_or_create(
                title=title[:300],
                author=author,
                defaults=dict(
                    category=category,
                    publisher=publisher[:200],
                    publication_year=year,
                    description=description,
                    price=price[:50],
                ),
            )
            if created:
                created_books += 1
            else:
                reused_books += 1

            inv_number = f"FOND-{old_id or (BookCopy.objects.count() + 1):>05}"
            if not BookCopy.objects.filter(inventory_number=inv_number).exists():
                BookCopy.objects.create(book=book, inventory_number=inv_number)
                created_copies += 1

        self.stdout.write("Nusxa sonlarini qayta hisoblash...")
        for book in Book.objects.all():
            book.recalc_quantities()

        self.stdout.write(self.style.SUCCESS(
            "\nImport tugadi!\n"
            f"  Jami qatorlar:             {total}\n"
            f"  Yangi yaratilgan kitoblar:  {created_books}\n"
            f"  Mavjud kitobga qo'shilgan:  {reused_books}\n"
            f"  Yaratilgan nusxalar:        {created_copies}\n"
            f"  O'tkazib yuborilgan:        {skipped}\n"
        ))
