"""8-modul: Hisobotlar - ekranda ko'rish va Excel/PDF ga eksport qilish."""
from datetime import timedelta

from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from django.views.generic import TemplateView, View

from apps.books.models import Book, BookCopy, CopyStatus
from apps.catalog.models import Category
from apps.core.mixins import LibrarianRequiredMixin
from apps.loans.models import Loan
from apps.readers.models import Reader, ReaderStatus

REPORT_TYPES = [
    ("fund", "Kutubxona fondi (kitoblar va nusxalar)"),
    ("loans", "Berilgan kitoblar"),
    ("overdue", "Muddati o'tgan kitoblar"),
    ("readers", "Kitobxonlar ro'yxati"),
    ("popular", "Eng ko'p o'qilgan kitoblar"),
]


def _period(request):
    """So'rovdan sana oralig'ini oladi (standart: so'nggi 30 kun)."""
    today = timezone.localdate()
    date_from = request.GET.get("from") or (today - timedelta(days=30)).isoformat()
    date_to = request.GET.get("to") or today.isoformat()
    return date_from, date_to


def build_report(report_type, date_from, date_to):
    """Hisobot uchun sarlavhalar va qatorlarni tayyorlaydi."""
    if report_type == "loans":
        rows = (
            Loan.objects.select_related("reader", "copy__book", "issued_by")
            .filter(issue_date__gte=date_from, issue_date__lte=date_to)
            .order_by("-issue_date")
        )
        headers = ["№", "Kitobxon", "Kodi", "Kitob", "Inventar",
                   "Berilgan", "Muddat", "Qaytarilgan", "Holat"]
        data = [
            [
                i,
                loan.reader.full_name,
                loan.reader.reader_code,
                loan.copy.book.title,
                loan.copy.inventory_number,
                loan.issue_date.strftime("%d.%m.%Y"),
                loan.due_date.strftime("%d.%m.%Y"),
                loan.return_date.strftime("%d.%m.%Y") if loan.return_date else "—",
                loan.get_status_display(),
            ]
            for i, loan in enumerate(rows, 1)
        ]
        return "Berilgan kitoblar hisoboti", headers, data

    if report_type == "overdue":
        today = timezone.localdate()
        rows = (
            Loan.objects.select_related("reader", "copy__book")
            .filter(return_date__isnull=True, due_date__lt=today)
            .order_by("due_date")
        )
        headers = ["№", "Kitobxon", "Telefon", "Kitob", "Inventar",
                   "Muddat", "Kechikish (kun)"]
        data = [
            [
                i,
                loan.reader.full_name,
                loan.reader.phone or "—",
                loan.copy.book.title,
                loan.copy.inventory_number,
                loan.due_date.strftime("%d.%m.%Y"),
                loan.overdue_days,
            ]
            for i, loan in enumerate(rows, 1)
        ]
        return "Muddati o'tgan kitoblar", headers, data

    if report_type == "readers":
        rows = Reader.objects.annotate(
            loans_total=Count("loans"),
            loans_active=Count("loans", filter=Q(loans__return_date__isnull=True)),
        ).order_by("full_name")
        headers = ["№", "Kodi", "F.I.Sh.", "Turi", "Telefon",
                   "Ro'yxatdan", "Holat", "Jami olgan", "Qo'lida"]
        data = [
            [
                i, r.reader_code, r.full_name, r.get_reader_type_display(),
                r.phone or "—", r.registration_date.strftime("%d.%m.%Y"),
                r.get_status_display(), r.loans_total, r.loans_active,
            ]
            for i, r in enumerate(rows, 1)
        ]
        return "Kitobxonlar ro'yxati", headers, data

    if report_type == "popular":
        rows = (
            Book.objects.select_related("author")
            .annotate(
                loans_total=Count(
                    "copies__loans",
                    filter=Q(copies__loans__issue_date__gte=date_from,
                             copies__loans__issue_date__lte=date_to),
                )
            )
            .filter(loans_total__gt=0)
            .order_by("-loans_total")[:100]
        )
        headers = ["№", "Kitob", "Muallif", "Toifa", "Berilgan marta", "Mavjud nusxa"]
        data = [
            [
                i, b.title, b.author.full_name, b.category.name,
                b.loans_total, b.available_quantity,
            ]
            for i, b in enumerate(rows, 1)
        ]
        return "Eng ko'p o'qilgan kitoblar", headers, data

    # Standart: fond hisoboti
    rows = (
        Book.objects.select_related("author", "category")
        .annotate(copies_total=Count("copies"))
        .order_by("category__name", "title")
    )
    headers = ["№", "Kitob", "Muallif", "Toifa", "ISBN", "Nashr yili",
               "Jami nusxa", "Mavjud", "Javon"]
    data = [
        [
            i, b.title, b.author.full_name, b.category.name, b.isbn or "—",
            b.publication_year or "—", b.total_quantity, b.available_quantity,
            b.shelf_location or "—",
        ]
        for i, b in enumerate(rows, 1)
    ]
    return "Kutubxona fondi hisoboti", headers, data


class ReportHomeView(LibrarianRequiredMixin, TemplateView):
    """Hisobotlar sahifasi: umumiy ko'rsatkichlar va hisobot tanlash."""

    template_name = "reports/report_home.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        date_from, date_to = _period(self.request)
        report_type = self.request.GET.get("type", "fund")

        title, headers, data = build_report(report_type, date_from, date_to)
        ctx.update(
            {
                "report_types": REPORT_TYPES,
                "selected_type": report_type,
                "date_from": date_from,
                "date_to": date_to,
                "report_title": title,
                "headers": headers,
                "rows": data[:300],
                "row_count": len(data),
                "truncated": len(data) > 300,
            }
        )

        # Yuqoridagi umumiy ko'rsatkichlar
        today = timezone.localdate()
        ctx["summary"] = {
            "titles": Book.objects.count(),
            "copies": BookCopy.objects.exclude(
                status__in=[CopyStatus.WRITTEN_OFF, CopyStatus.LOST]
            ).count(),
            "readers": Reader.objects.filter(status=ReaderStatus.ACTIVE).count(),
            "issued": Loan.objects.filter(return_date__isnull=True).count(),
            "overdue": Loan.objects.filter(
                return_date__isnull=True, due_date__lt=today
            ).count(),
            "period_loans": Loan.objects.filter(
                issue_date__gte=date_from, issue_date__lte=date_to
            ).count(),
        }
        ctx["by_category"] = (
            Category.objects.annotate(books_total=Count("books"))
            .filter(books_total__gt=0)
            .order_by("-books_total")[:10]
        )
        return ctx


class ExportExcelView(LibrarianRequiredMixin, View):
    """Hisobotni .xlsx fayl sifatida yuklab berish."""

    def get(self, request):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            return HttpResponse(
                "Excel eksport uchun openpyxl kerak: pip install openpyxl",
                status=500,
            )

        report_type = request.GET.get("type", "fund")
        date_from, date_to = _period(request)
        title, headers, data = build_report(report_type, date_from, date_to)

        wb = Workbook()
        ws = wb.active
        ws.title = "Hisobot"

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        cell = ws.cell(row=1, column=1, value=title)
        cell.font = Font(size=14, bold=True, color="1E3A5F")
        cell.alignment = Alignment(horizontal="center")

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        sub = ws.cell(
            row=2, column=1,
            value=f"Davr: {date_from} — {date_to} | "
                  f"Shakllantirildi: {timezone.localtime():%d.%m.%Y %H:%M}",
        )
        sub.alignment = Alignment(horizontal="center")

        head_fill = PatternFill("solid", fgColor="1E3A5F")
        thin = Side(style="thin", color="D0D7E2")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, name in enumerate(headers, 1):
            c = ws.cell(row=4, column=col, value=name)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border

        for r, row in enumerate(data, 5):
            for col, value in enumerate(row, 1):
                c = ws.cell(row=r, column=col, value=value)
                c.border = border
                if col == 1:
                    c.alignment = Alignment(horizontal="center")

        for col, name in enumerate(headers, 1):
            longest = max(
                [len(str(name))] + [len(str(row[col - 1])) for row in data[:200]] or [10]
            )
            ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 3, 10), 45)
        ws.freeze_panes = "A5"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument."
                         "spreadsheetml.sheet"
        )
        filename = f"hisobot_{report_type}_{timezone.localdate():%Y%m%d}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


class ExportPdfView(LibrarianRequiredMixin, View):
    """Hisobotni .pdf fayl sifatida yuklab berish."""

    def get(self, request):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
            )
        except ImportError:
            return HttpResponse(
                "PDF eksport uchun reportlab kerak: pip install reportlab",
                status=500,
            )

        report_type = request.GET.get("type", "fund")
        date_from, date_to = _period(request)
        title, headers, data = build_report(report_type, date_from, date_to)
        data = data[:1000]

        response = HttpResponse(content_type="application/pdf")
        filename = f"hisobot_{report_type}_{timezone.localdate():%Y%m%d}.pdf"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

        doc = SimpleDocTemplate(
            response, pagesize=landscape(A4),
            leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm,
            title=title,
        )
        styles = getSampleStyleSheet()
        cell_style = ParagraphStyle(
            "cell", parent=styles["Normal"], fontSize=7.5, leading=9.5
        )
        head_style = ParagraphStyle(
            "head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold"
        )

        story = [
            Paragraph(f"<b>{title}</b>", styles["Title"]),
            Paragraph(
                f"Davr: {date_from} — {date_to} &nbsp;|&nbsp; Jami: {len(data)} ta yozuv "
                f"&nbsp;|&nbsp; {timezone.localtime():%d.%m.%Y %H:%M}",
                styles["Normal"],
            ),
            Spacer(1, 6 * mm),
        ]

        table_data = [[Paragraph(str(h), head_style) for h in headers]]
        for row in data:
            table_data.append([Paragraph(str(v), cell_style) for v in row])

        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C8D2E0")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                     [colors.white, colors.HexColor("#F3F6FA")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(table)
        doc.build(story)
        return response
